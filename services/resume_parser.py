
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


MAX_RESUME_BYTES = 10 * 1024 * 1024  # 10 MB
MAX_RESUME_CHARACTERS = 60_000


def extract_resume_text(
    file_bytes: bytes,
    filename: str,
) -> str:
    """Extract readable text from a PDF or DOCX resume."""

    if not file_bytes:
        raise ValueError("The uploaded resume is empty.")

    if len(file_bytes) > MAX_RESUME_BYTES:
        raise ValueError("Resume exceeds the 10 MB upload limit.")

    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        text = _extract_pdf(file_bytes)
    elif extension == ".docx":
        text = _extract_docx(file_bytes)
    else:
        raise ValueError("Upload a PDF or DOCX resume.")

    text = text.strip()

    if not text:
        raise ValueError(
            "No readable text was found. If this is a scanned PDF, "
            "run OCR on it first."
        )

    return text[:MAX_RESUME_CHARACTERS]


def _extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))

    if reader.is_encrypted:
        raise ValueError(
            "This PDF is password-protected. Upload an unlocked copy."
        )

    pages = []

    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            pages.append(page_text)

    return "\n\n".join(pages)


def _extract_docx(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))

    parts = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    # Resume skills and work history are sometimes stored in tables.
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)
